"""plans/DATA_POPULATION_PLAN.md — tools/bulk_import.py, the standalone bulk loader.

Covers the four things the loader is specifically for:
  - streaming parse (never materialises the whole sheet)
  - validation identical to the existing UI import path (reuses validate_import_rows)
  - --dry-run writes nothing
  - resume after a hard kill mid-batch: no duplicated rows, no skipped rows

Every DB-touching test uses an in-memory fake connection (never a real database),
matching the pattern in tests/test_mobile_lookup_pdf_report.py.
"""
import inspect
import itertools
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import tools.bulk_import as bulk_import  # noqa: E402


# ─────────────────────────────────────────────
# Streaming parse
# ─────────────────────────────────────────────
class _CountingRowSource:
    """Models a worksheet with `total_rows` data rows, one header row.

    `pulled` only advances as the consumer actually calls next() on the generator this
    produces — i.e. it's pull-based, exactly like openpyxl's read_only worksheet.
    """

    def __init__(self, total_rows):
        self.total_rows = total_rows
        self.pulled = 0

    def __iter__(self):
        yield ("h1", "h2")
        for i in range(self.total_rows):
            self.pulled += 1
            yield (i, i * 2)


class _FakeWorksheet:
    def __init__(self, source):
        self._source = source

    def iter_rows(self, values_only=True):
        assert values_only is True
        return iter(self._source)


class StreamingParseTests(unittest.TestCase):
    def test_rows_from_worksheet_is_a_generator(self):
        ws = _FakeWorksheet(_CountingRowSource(10))
        gen = bulk_import._rows_from_worksheet(ws)
        self.assertTrue(inspect.isgenerator(gen))

    def test_never_materialises_the_whole_sheet(self):
        """The core streaming claim: pulling a handful of rows out of a million-row
        source must not touch anywhere near a million rows of the underlying source.
        A `list(ws.iter_rows(...))` implementation would pull all 1,000,000 rows
        immediately; a streaming one pulls only what's consumed.
        """
        source = _CountingRowSource(1_000_000)
        ws = _FakeWorksheet(source)
        gen = bulk_import._rows_from_worksheet(ws)

        first = next(gen)
        self.assertEqual(first, {"h1": 0, "h2": 0})
        self.assertEqual(source.pulled, 1)

        for _ in range(4):
            next(gen)
        self.assertEqual(source.pulled, 5)
        self.assertLess(source.pulled, 1_000_000)

    def test_chunked_only_holds_one_batch_at_a_time(self):
        source = _CountingRowSource(1_000_000)
        ws = _FakeWorksheet(source)
        gen = bulk_import._rows_from_worksheet(ws)
        chunks = bulk_import._chunked(gen, 10)
        first_chunk = next(chunks)
        self.assertEqual(len(first_chunk), 10)
        self.assertEqual(source.pulled, 10)  # not 1,000,000

    def test_empty_sheet_yields_nothing(self):
        ws = _FakeWorksheet(_CountingRowSource(0))
        # Only the header row exists once total_rows=0; _rows_from_worksheet should
        # yield zero data rows.
        rows = list(bulk_import._rows_from_worksheet(ws))
        self.assertEqual(rows, [])

    def test_truly_empty_source_returns_without_error(self):
        class _Empty:
            def iter_rows(self, values_only=True):
                return iter([])

        rows = list(bulk_import._rows_from_worksheet(_Empty()))
        self.assertEqual(rows, [])


# ─────────────────────────────────────────────
# iter_xlsx_rows() vs the existing app._parse_import_file() — identical output,
# identical validate_import_rows() results.
# ─────────────────────────────────────────────
class _FakeFileStorage:
    """Minimal stand-in for a Werkzeug FileStorage: .filename + .read()."""

    def __init__(self, filename, data):
        self.filename = filename
        self._data = data

    def read(self):
        return self._data


class ParityWithExistingParserTests(unittest.TestCase):
    def setUp(self):
        import openpyxl

        self.tmpdir = self._make_tmpdir()
        self.path = os.path.join(self.tmpdir, "subscribers.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["mobile_number", "subscriber_name", "district"])
        ws.append(["9876543210", "Ramesh Kumar", "Pune"])
        ws.append(["9876543211", "Suresh Rao", "Nagpur"])
        ws.append(["9876543210", "Duplicate Number", "Pune"])  # duplicate key
        ws.append(["", "Missing Number", "Nowhere"])  # missing required field
        wb.save(self.path)

    def _make_tmpdir(self):
        import shutil
        import tempfile

        d = tempfile.mkdtemp(prefix="bulk_import_test_")
        self.addCleanup(shutil.rmtree, d, ignore_errors=True)
        return d

    def test_streamed_rows_match_existing_parser_output(self):
        import app as app_module

        with open(self.path, "rb") as fh:
            raw = fh.read()
        rows_existing = app_module._parse_import_file(_FakeFileStorage("subscribers.xlsx", raw))
        rows_streamed = list(bulk_import.iter_xlsx_rows(self.path))
        self.assertEqual(rows_existing, rows_streamed)

    def test_validate_import_rows_gives_identical_results_on_both_paths(self):
        import app as app_module

        with open(self.path, "rb") as fh:
            raw = fh.read()
        rows_existing = app_module._parse_import_file(_FakeFileStorage("subscribers.xlsx", raw))
        rows_streamed = list(bulk_import.iter_xlsx_rows(self.path))

        out_a, issues_a, summary_a = app_module.validate_import_rows("subscriber", rows_existing, {})
        out_b, issues_b, summary_b = app_module.validate_import_rows("subscriber", rows_streamed, {})

        # validate_import_rows(module, rows, {}) — an empty mapping, exactly what the
        # CLI always passes — sets row['_raw'] = row (the same object, not a copy),
        # since with no mapping `row = raw` is done first. That makes every processed
        # row self-referential; comparing them with == recurses forever (confirmed:
        # this test hit RecursionError before this strip). Strip '_raw' before
        # comparing — this is a test-harness workaround for that latent behaviour, not
        # a claim that '_raw' differs between the two paths (see _raw_payload() in
        # tools/bulk_import.py for how the loader itself avoids the same trap).
        def _without_raw(rows):
            return [{k: v for k, v in row.items() if k != "_raw"} for row in rows]

        self.assertEqual(_without_raw(out_a), _without_raw(out_b))
        self.assertEqual(issues_a, issues_b)
        self.assertEqual(summary_a, summary_b)
        # Sanity: the fixture actually exercises both a duplicate and a missing-field row.
        self.assertEqual(summary_a["duplicate_count"], 1)
        self.assertEqual(summary_a["error_count"], 1)


# ─────────────────────────────────────────────
# Fake DB for the write/resume tests
# ─────────────────────────────────────────────
class FakeCursor:
    def __init__(self, conn, dictionary=False):
        self.conn = conn
        self.dictionary = dictionary
        self.lastrowid = None
        self._result = []

    def execute(self, query, params=()):
        q = " ".join(query.split())
        self._result = []

        if q.startswith("INSERT INTO import_batches"):
            row_id = self.conn._next_id("import_batches")
            row = {
                "id": row_id,
                "source_type": params[0],
                "original_filename": params[1],
                "total_rows": params[2],
                "imported_by": params[3],
                "status": "in_progress",
                "rows_committed": 0,
                "source_file": params[4],
                "file_checksum": params[5],
            }
            self.conn._stage(lambda tables, row=row: tables["import_batches"].append(row))
            self.lastrowid = row_id

        elif q.startswith("SELECT * FROM import_batches WHERE source_type=%s AND source_file=%s AND file_checksum=%s"):
            src_type, src_file, checksum = params
            matches = [
                r for r in self.conn.tables["import_batches"]
                if r["source_type"] == src_type and r["source_file"] == src_file and r["file_checksum"] == checksum
            ]
            matches.sort(key=lambda r: -r["id"])
            self._result = matches[:1]

        elif q.startswith("UPDATE import_batches SET rows_committed=%s WHERE id=%s"):
            rc, bid = params

            def op(tables, rc=rc, bid=bid):
                for r in tables["import_batches"]:
                    if r["id"] == bid:
                        r["rows_committed"] = rc

            self.conn._stage(op)

        elif q.startswith("UPDATE import_batches SET status=%s, total_rows=%s WHERE id=%s"):
            status, total, bid = params

            def op(tables, status=status, total=total, bid=bid):
                for r in tables["import_batches"]:
                    if r["id"] == bid:
                        r["status"] = status
                        r["total_rows"] = total

            self.conn._stage(op)

        elif q.startswith("SELECT * FROM bts_tower_records WHERE lac=%s AND cell_id=%s"):
            self._result = []

        else:
            raise AssertionError(f"FakeCursor got an unexpected execute: {q!r}")

    def executemany(self, query, seq_params):
        q = " ".join(query.split())
        rows = list(seq_params)

        if q.startswith("INSERT INTO cdr_records"):
            def op(tables, rows=rows):
                for p in rows:
                    row_id = self.conn._next_id("cdr_records")
                    tables["cdr_records"].append({"id": row_id, "params": p})

            self.conn._stage(op)

        elif q.startswith("INSERT INTO subscriber_records"):
            def op(tables, rows=rows):
                for p in rows:
                    row_id = self.conn._next_id("subscriber_records")
                    tables["subscriber_records"].append({"id": row_id, "params": p})

            self.conn._stage(op)

        else:
            raise AssertionError(f"FakeCursor got an unexpected executemany: {q!r}")

    def fetchone(self):
        return self._result[0] if self._result else None

    def fetchall(self):
        return self._result


class FakeConn:
    """Models MySQL transaction semantics well enough for the resume test: writes made
    since the last commit() are only applied to `tables` (the "durable" state) when
    commit() succeeds. If crash_after_n_commits triggers, the pending writes for that
    commit are discarded instead of applied — modelling a hard kill / connection drop
    before the transaction reached the server, i.e. exactly what a real InnoDB
    connection does with anything that was never committed.
    """

    def __init__(self, tables=None, id_counters=None, crash_after_n_commits=None):
        self.tables = tables if tables is not None else {"import_batches": [], "cdr_records": [], "subscriber_records": []}
        self._id_counters = id_counters if id_counters is not None else {}
        self._pending = []
        self.commit_count = 0
        self.crash_after_n_commits = crash_after_n_commits

    def _next_id(self, table):
        self._id_counters[table] = self._id_counters.get(table, 0) + 1
        return self._id_counters[table]

    def _stage(self, op):
        self._pending.append(op)

    def cursor(self, dictionary=False):
        return FakeCursor(self, dictionary=dictionary)

    def commit(self):
        self.commit_count += 1
        if self.crash_after_n_commits is not None and self.commit_count == self.crash_after_n_commits:
            self._pending = []
            raise RuntimeError("simulated hard kill mid-commit")
        for op in self._pending:
            op(self.tables)
        self._pending = []

    def rollback(self):
        self._pending = []

    def close(self):
        pass


class _ExplodingConn:
    """Fails the test the moment dry-run touches the DB connection at all —
    not just writes: no cursor(), no execute(), no commit(). A real dry-run needs
    no DB dependency beyond what the caller already opened.
    """

    def cursor(self, dictionary=False):
        raise AssertionError("dry-run must not open a cursor")

    def commit(self):
        raise AssertionError("dry-run must not commit")

    def close(self):
        pass


def _write_cdr_xlsx(path, n_rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["calling_number", "called_number", "start_time"])
    for i in range(n_rows):
        ws.append([f"90000{i:05d}", f"80000{i:05d}", "2026-08-01 10:00:00"])
    wb.save(path)


def _write_subscriber_xlsx(path, n_rows):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["mobile_number", "subscriber_name"])
    for i in range(n_rows):
        ws.append([f"90000{i:05d}", f"Person {i}"])
    wb.save(path)


class DryRunTests(unittest.TestCase):
    def setUp(self):
        import shutil
        import tempfile

        self.tmpdir = tempfile.mkdtemp(prefix="bulk_import_dryrun_")
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.path = os.path.join(self.tmpdir, "subs.xlsx")
        _write_subscriber_xlsx(self.path, 12)

    def test_dry_run_writes_nothing_and_reports_counts(self):
        import app as app_module

        conn = _ExplodingConn()
        result = bulk_import.import_file(
            app_module, conn, "subscriber", self.path,
            batch_size=5, case_ref=None, resume=False, dry_run=True, log=lambda *a: None,
        )
        self.assertTrue(result["dry_run"])
        self.assertEqual(result["rows_committed"], 12)
        self.assertEqual(result["rows_seen"], 12)
        self.assertEqual(result["error_rows"], 0)
        self.assertIsNone(result["batch_id"])


class ResumeAfterHardKillTests(unittest.TestCase):
    """The requirement this exists to prove: an interrupted run does not duplicate
    committed rows and does not skip unprocessed ones, even when the interruption
    lands mid-batch (rows already executemany()'d but never committed).
    """

    def setUp(self):
        import shutil
        import tempfile

        self.tmpdir = tempfile.mkdtemp(prefix="bulk_import_resume_")
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)
        self.path = os.path.join(self.tmpdir, "cdr.xlsx")
        self.n_rows = 23
        _write_cdr_xlsx(self.path, self.n_rows)

    def test_crash_mid_batch_then_resume_yields_exactly_one_copy_of_every_row(self):
        import app as app_module

        tables = {"import_batches": [], "cdr_records": [], "subscriber_records": []}
        id_counters = {}

        # commit #1 = creating the new import_batches row.
        # commit #2 = the first batch (rows 1-5) — succeeds.
        # commit #3 = the second batch (rows 6-10) — this is where we simulate the
        # hard kill: executemany() has already staged the INSERT for rows 6-10 *and*
        # the rows_committed=10 update, but commit() raises before either is applied.
        conn1 = FakeConn(tables=tables, id_counters=id_counters, crash_after_n_commits=3)
        with self.assertRaises(RuntimeError):
            bulk_import.import_file(
                app_module, conn1, "cdr", self.path,
                batch_size=5, case_ref="FIR-2026-1", resume=False, dry_run=False,
                log=lambda *a: None,
            )

        # Only the first batch is durable: 5 cdr rows, rows_committed=5, still in_progress.
        self.assertEqual(len(tables["cdr_records"]), 5)
        batch = tables["import_batches"][0]
        self.assertEqual(batch["rows_committed"], 5)
        self.assertEqual(batch["status"], "in_progress")

        # "process restart": a fresh connection over the same durable tables/id counters.
        conn2 = FakeConn(tables=tables, id_counters=id_counters, crash_after_n_commits=None)
        result = bulk_import.import_file(
            app_module, conn2, "cdr", self.path,
            batch_size=5, case_ref="FIR-2026-1", resume=True, dry_run=False,
            log=lambda *a: None,
        )

        self.assertEqual(result["rows_skipped_via_resume"], 5)
        self.assertEqual(result["rows_committed"], self.n_rows)
        self.assertEqual(result["rows_seen"], self.n_rows)

        # No duplicates, nothing missing: exactly n_rows cdr_records rows, each for a
        # distinct (calling_number, called_number) pair.
        self.assertEqual(len(tables["cdr_records"]), self.n_rows)
        pairs = [(p["params"][1], p["params"][2]) for p in tables["cdr_records"]]
        self.assertEqual(len(pairs), len(set(pairs)))

        final_batch = tables["import_batches"][0]
        self.assertEqual(final_batch["rows_committed"], self.n_rows)
        self.assertEqual(final_batch["status"], "imported")
        self.assertEqual(final_batch["total_rows"], self.n_rows)

    def test_resume_without_flag_refuses_to_start_a_duplicate_load(self):
        import app as app_module

        tables = {"import_batches": [], "cdr_records": [], "subscriber_records": []}
        id_counters = {}

        conn1 = FakeConn(tables=tables, id_counters=id_counters, crash_after_n_commits=3)
        with self.assertRaises(RuntimeError):
            bulk_import.import_file(
                app_module, conn1, "cdr", self.path,
                batch_size=5, case_ref=None, resume=False, dry_run=False, log=lambda *a: None,
            )

        conn2 = FakeConn(tables=tables, id_counters=id_counters)
        with self.assertRaises(RuntimeError) as ctx:
            bulk_import.import_file(
                app_module, conn2, "cdr", self.path,
                batch_size=5, case_ref=None, resume=False, dry_run=False, log=lambda *a: None,
            )
        self.assertIn("--resume", str(ctx.exception))
        # Nothing new was written by the refused second attempt.
        self.assertEqual(len(tables["cdr_records"]), 5)

    def test_clean_run_with_no_prior_batch_commits_every_row_exactly_once(self):
        import app as app_module

        tables = {"import_batches": [], "cdr_records": [], "subscriber_records": []}
        conn = FakeConn(tables=tables)
        result = bulk_import.import_file(
            app_module, conn, "cdr", self.path,
            batch_size=7, case_ref=None, resume=False, dry_run=False, log=lambda *a: None,
        )
        self.assertEqual(result["rows_committed"], self.n_rows)
        self.assertEqual(len(tables["cdr_records"]), self.n_rows)
        self.assertEqual(tables["import_batches"][0]["status"], "imported")


class FileFingerprintTests(unittest.TestCase):
    def setUp(self):
        import shutil
        import tempfile

        self.tmpdir = tempfile.mkdtemp(prefix="bulk_import_fp_")
        self.addCleanup(shutil.rmtree, self.tmpdir, ignore_errors=True)

    def test_same_file_same_fingerprint(self):
        path = os.path.join(self.tmpdir, "a.xlsx")
        with open(path, "wb") as fh:
            fh.write(b"hello")
        self.assertEqual(bulk_import.file_fingerprint(path), bulk_import.file_fingerprint(path))

    def test_different_size_different_fingerprint(self):
        path_a = os.path.join(self.tmpdir, "a.xlsx")
        path_b = os.path.join(self.tmpdir, "b.xlsx")
        with open(path_a, "wb") as fh:
            fh.write(b"hello")
        with open(path_b, "wb") as fh:
            fh.write(b"hello world, much longer")
        self.assertNotEqual(bulk_import.file_fingerprint(path_a), bulk_import.file_fingerprint(path_b))


if __name__ == "__main__":
    unittest.main()
